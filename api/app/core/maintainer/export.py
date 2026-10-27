"""Exportación de un listado a Excel (.xlsx)."""

from collections.abc import Iterable, Sequence
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO
from typing import Any
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from app.core.maintainer.spec import ExportColumn

XLSX_MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

_HEADER_FONT = Font(bold=True, color="FFFFFF")
_HEADER_FILL = PatternFill("solid", fgColor="324664")


def _cell_value(value: Any, tz: ZoneInfo) -> Any:
    if isinstance(value, bool):
        return "Sí" if value else "No"
    if isinstance(value, datetime):
        # Excel no maneja zonas: se escribe la hora local de la planta, sin zona.
        local = value.astimezone(tz) if value.tzinfo else value
        return local.replace(tzinfo=None)
    if isinstance(value, Decimal | int | float | str | date) or value is None:
        return value
    return str(value)


def _extract(obj: Any, column: ExportColumn) -> Any:
    return getattr(obj, column.value) if isinstance(column.value, str) else column.value(obj)


def build_workbook(
    rows: Iterable[Any], columns: Sequence[ExportColumn], *, sheet_title: str, timezone: str
) -> bytes:
    tz = ZoneInfo(timezone)
    wb = Workbook(write_only=True)
    ws = wb.create_sheet(title=sheet_title[:31])
    for index, column in enumerate(columns, start=1):
        ws.column_dimensions[get_column_letter(index)].width = column.width
    ws.freeze_panes = "A2"

    header = []
    for column in columns:
        cell = WriteOnlyCell(ws, value=column.header)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        header.append(cell)
    ws.append(header)

    for obj in rows:
        values = []
        for column in columns:
            value = _cell_value(_extract(obj, column), tz)
            cell = WriteOnlyCell(ws, value=value)
            if isinstance(value, datetime):
                cell.number_format = "dd-mm-yyyy hh:mm"
            elif isinstance(value, date):
                cell.number_format = "dd-mm-yyyy"
            values.append(cell)
        ws.append(values)

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
