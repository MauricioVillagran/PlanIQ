"""Patrón de mantenedor, probado con el mantenedor de plantas (implementación de referencia)."""

from collections.abc import Callable
from io import BytesIO
from typing import Any

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.modules.plants.models import Plant
from app.modules.users.models import User
from tests.integration.conftest import auth_header, login

pytestmark = pytest.mark.integration

URL = "/api/v1/plants"


@pytest.fixture
def admin(client: TestClient, make_user: Callable[..., User]) -> dict[str, str]:
    make_user("jefa", roles=("admin",))
    return auth_header(login(client, "jefa")["access_token"])


def _create(client: TestClient, headers: dict[str, str], **overrides: Any) -> dict[str, Any]:
    body = {"code": "VS", "name": "Planta Victoria Sur", "plant_type": "own"} | overrides
    res = client.post(URL, json=body, headers=headers)
    assert res.status_code == 201, res.text
    data: dict[str, Any] = res.json()
    return data


def test_create_returns_record_with_row_version(client: TestClient, admin: dict[str, str]) -> None:
    plant = _create(client, admin, location="  Puerto Montt  ")
    assert plant["code"] == "VS"
    assert plant["location"] == "Puerto Montt"  # recorta espacios
    assert plant["timezone"] == "America/Santiago"
    assert plant["is_active"] is True
    assert len(plant["row_version"]) == 16


@pytest.mark.parametrize(
    ("body", "field", "code"),
    [
        ({"code": "X", "name": "X", "plant_type": "propia"}, "plant_type", "literal_error"),
        (
            {"code": "X", "name": "X", "plant_type": "own", "timezone": "Chile/Marte"},
            "timezone",
            "value_error",
        ),
        ({"code": "X", "name": "X", "plant_type": "own", "color": "rojo"}, "color", "extra_forbidden"),
        ({"name": "X", "plant_type": "own"}, "code", "missing"),
    ],
)
def test_create_validation_errors(
    client: TestClient, admin: dict[str, str], body: dict[str, Any], field: str, code: str
) -> None:
    res = client.post(URL, json=body, headers=admin)
    assert res.status_code == 422
    errors = {e["field"]: e["code"] for e in res.json()["field_errors"]}
    assert errors[field] == code


def test_duplicate_code_is_a_field_error(client: TestClient, admin: dict[str, str]) -> None:
    _create(client, admin)
    res = client.post(URL, json={"code": "VS", "name": "Otra", "plant_type": "own"}, headers=admin)
    assert res.status_code == 422
    body = res.json()
    assert body["code"] == "duplicate"
    assert body["field_errors"] == [
        {"field": "code", "message": "Ya existe planta con ese código.", "code": "unique"}
    ]


def test_update_with_current_row_version(client: TestClient, admin: dict[str, str]) -> None:
    plant = _create(client, admin)
    res = client.patch(
        f"{URL}/{plant['id']}",
        json={"name": "Victoria Sur", "row_version": plant["row_version"]},
        headers=admin,
    )
    assert res.status_code == 200
    updated = res.json()
    assert updated["name"] == "Victoria Sur"
    assert updated["code"] == "VS"  # campos no enviados no cambian
    assert updated["row_version"] != plant["row_version"]
    assert updated["updated_at"] is not None


def test_update_with_stale_row_version_returns_409_with_current_data(
    client: TestClient, admin: dict[str, str]
) -> None:
    plant = _create(client, admin)
    first = client.patch(
        f"{URL}/{plant['id']}", json={"name": "Cambio A", "row_version": plant["row_version"]}, headers=admin
    )
    assert first.status_code == 200

    # Segundo usuario editando con la versión que leyó antes del primer cambio.
    second = client.patch(
        f"{URL}/{plant['id']}", json={"name": "Cambio B", "row_version": plant["row_version"]}, headers=admin
    )
    assert second.status_code == 409
    body = second.json()
    assert body["code"] == "conflict"
    assert body["details"]["current"]["name"] == "Cambio A"
    assert body["details"]["current"]["row_version"] == first.json()["row_version"]


def test_update_requires_row_version(client: TestClient, admin: dict[str, str]) -> None:
    plant = _create(client, admin)
    res = client.patch(f"{URL}/{plant['id']}", json={"name": "Sin versión"}, headers=admin)
    assert res.status_code == 422
    assert res.json()["field_errors"][0]["field"] == "row_version"


def test_patch_cannot_null_a_required_field(client: TestClient, admin: dict[str, str]) -> None:
    plant = _create(client, admin)
    res = client.patch(
        f"{URL}/{plant['id']}", json={"name": None, "row_version": plant["row_version"]}, headers=admin
    )
    assert res.status_code == 422
    assert res.json()["code"] == "required"
    assert res.json()["field_errors"][0]["field"] == "name"


def test_deactivate_and_activate(client: TestClient, admin: dict[str, str]) -> None:
    plant = _create(client, admin)
    off = client.post(f"{URL}/{plant['id']}/deactivate", headers=admin)
    assert off.status_code == 200 and off.json()["is_active"] is False
    # Idempotente: desactivar de nuevo no falla.
    assert client.post(f"{URL}/{plant['id']}/deactivate", headers=admin).status_code == 200

    active = client.get(URL, params={"status": "active"}, headers=admin).json()
    inactive = client.get(URL, params={"status": "inactive"}, headers=admin).json()
    assert plant["id"] not in [p["id"] for p in active["items"]]
    assert plant["id"] in [p["id"] for p in inactive["items"]]

    on = client.post(f"{URL}/{plant['id']}/activate", headers=admin)
    assert on.json()["is_active"] is True


def test_list_search_filter_sort_and_pagination(client: TestClient, admin: dict[str, str]) -> None:
    _create(client, admin, code="A1", name="Planta Alfa", location="Calbuco")
    _create(client, admin, code="B2", name="Planta Beta", plant_type="contract")
    _create(client, admin, code="C3", name="Gamma", location="Puerto Montt")

    by_text = client.get(URL, params={"q": "planta"}, headers=admin).json()
    assert {p["code"] for p in by_text["items"]} >= {"A1", "B2"}
    assert "C3" not in {p["code"] for p in by_text["items"]}

    by_location = client.get(URL, params={"q": "calbu"}, headers=admin).json()
    assert [p["code"] for p in by_location["items"]] == ["A1"]

    contract = client.get(URL, params={"plant_type": "contract"}, headers=admin).json()
    assert [p["code"] for p in contract["items"]] == ["B2"]

    ordered = client.get(URL, params={"sort": "-code", "q": "3"}, headers=admin).json()
    assert ordered["items"][0]["code"] == "C3"

    page = client.get(URL, params={"page": 1, "page_size": 2, "sort": "code"}, headers=admin).json()
    assert len(page["items"]) == 2
    assert page["total"] >= 3


def test_search_escapes_like_wildcards(client: TestClient, admin: dict[str, str]) -> None:
    _create(client, admin, code="P1", name="Planta 100%")
    _create(client, admin, code="P2", name="Planta 1000")
    res = client.get(URL, params={"q": "100%"}, headers=admin).json()
    assert [p["code"] for p in res["items"]] == ["P1"]


def test_invalid_sort_field(client: TestClient, admin: dict[str, str]) -> None:
    res = client.get(URL, params={"sort": "password"}, headers=admin)
    assert res.status_code == 422
    assert res.json()["code"] == "invalid_sort"


def test_not_found(client: TestClient, admin: dict[str, str]) -> None:
    res = client.get(f"{URL}/999999", headers=admin)
    assert res.status_code == 404
    assert res.json()["message"] == "No existe planta con id 999999."


def test_permissions_view_and_edit(client: TestClient, make_user: Callable[..., User]) -> None:
    make_user("ana", roles=("planning",))
    make_user("sinrol")
    planner = auth_header(login(client, "ana")["access_token"])
    nobody = auth_header(login(client, "sinrol")["access_token"])

    assert client.get(URL, headers=planner).status_code == 200  # masters.view
    denied = client.post(URL, json={"code": "X", "name": "X", "plant_type": "own"}, headers=planner)
    assert denied.status_code == 403
    assert denied.json()["details"] == {"missing_permissions": ["masters.edit"]}
    assert client.get(URL, headers=nobody).status_code == 403


def test_users_only_see_their_plants(
    client: TestClient, make_user: Callable[..., User], make_plant: Callable[..., Plant]
) -> None:
    own, other = make_plant("P1"), make_plant("P2")
    make_user("ana", roles=("production",), plants=(own,))
    headers = auth_header(login(client, "ana")["access_token"])

    listed = client.get(URL, headers=headers).json()
    assert [p["code"] for p in listed["items"]] == ["P1"]
    assert client.get(f"{URL}/{own.id}", headers=headers).status_code == 200
    assert client.get(f"{URL}/{other.id}", headers=headers).status_code == 404


def test_history_lists_audited_changes(client: TestClient, admin: dict[str, str]) -> None:
    plant = _create(client, admin)
    client.patch(
        f"{URL}/{plant['id']}", json={"name": "Nuevo", "row_version": plant["row_version"]}, headers=admin
    )
    client.post(f"{URL}/{plant['id']}/deactivate", headers=admin)

    history = client.get(f"{URL}/{plant['id']}/history", headers=admin).json()
    assert [e["action"] for e in history["items"]] == ["deactivate", "update", "create"]
    assert history["items"][1]["changes"]["name"] == {"old": "Planta Victoria Sur", "new": "Nuevo"}
    assert all(e["username"] == "jefa" for e in history["items"])


def test_export_xlsx_uses_list_filters(client: TestClient, admin: dict[str, str]) -> None:
    _create(client, admin, code="A1", name="Planta Alfa")
    _create(client, admin, code="B2", name="Planta Beta", plant_type="contract")

    res = client.get(f"{URL}/export.xlsx", params={"plant_type": "contract"}, headers=admin)
    assert res.status_code == 200
    assert res.headers["content-type"].startswith("application/vnd.openxmlformats")
    assert 'filename="plants-' in res.headers["content-disposition"]

    ws = load_workbook(BytesIO(res.content)).active
    assert ws is not None
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0][:4] == ("Código", "Nombre", "Tipo", "Ubicación")
    assert rows[1][:3] == ("B2", "Planta Beta", "Maquila")
    assert rows[1][5] == "Sí"
    assert len(rows) == 2


def test_writes_are_audited_with_user(client: TestClient, admin: dict[str, str], db_session: Any) -> None:
    plant = _create(client, admin)
    db_obj = db_session.get(Plant, plant["id"])
    assert db_obj.created_by is not None
